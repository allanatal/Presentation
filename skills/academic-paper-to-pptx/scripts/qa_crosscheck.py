#!/usr/bin/env python3
"""
Deterministic QA cross-check for generated decks (shared by academic-paper-to-pptx
and pptx-to-pptx). No model judgment — verifies mechanically what it can and emits
explicit MANUAL lines for the rest. See references/qa-checklist.md.

Usage:
  paper mode (deck built from a paper; claims file per qa-checklist.md schema):
    python qa_crosscheck.py deck.pptx --claims deck_claims.json --mode paper \
        [--source-text paper_text.md] [--output deck_QA.md] [--notes]

  pptx mode (deck restyled from a source deck; claims = parse_pptx.py JSON):
    python qa_crosscheck.py deck.pptx --claims parsed.json --mode pptx \
        [--output deck_QA.md] [--notes]

Exit code: 0 if no ❌ findings, 1 otherwise (so build loops fail fast).
"""

import argparse
import datetime
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# ── Font-size floors (pt) — hard minimums from style-spec.md ──
FLOOR_BODY = 18
FLOOR_SUBBULLET = 16
FLOOR_TABLE = 14
FLOOR_DIAGRAM = 12
FLOOR_CITATION = 10
FLOOR_BADGE = 11

# Badge/citation are recognized by their canonical positions (style-spec.md)
BADGE_LEFT_MIN = 9_500_000    # EMU
BADGE_TOP_MAX = 600_000       # EMU
CITATION_LEFT_MIN = 9_500_000
CITATION_TOP_MIN = 6_200_000

FULL_SLIDE_IMAGE_FRACTION = 0.8  # picture covering >80% of slide area = screenshot

NUM_RE = re.compile(r"(?<![A-Za-z0-9_])(?:NCT\d{8}|\d+(?:,\d{3})*(?:\.\d+)?)")


def norm_text(s):
    """Normalize dashes/decimal marks so token comparison is robust."""
    if not s:
        return ""
    return (
        s.replace("–", "-").replace("—", "-").replace("−", "-")
        .replace("·", ".")
    )


def numeric_tokens(s):
    """All numeric tokens in a string (commas stripped, % dropped for matching)."""
    return [m.group(0).replace(",", "") for m in NUM_RE.finditer(norm_text(s))]


# ═══════════════════════════════════════════════════
#  Deck extraction
# ═══════════════════════════════════════════════════

def is_badge(shape):
    return (shape.left or 0) >= BADGE_LEFT_MIN and (shape.top or 0) <= BADGE_TOP_MAX


def is_citation(shape):
    return (shape.left or 0) >= CITATION_LEFT_MIN and (shape.top or 0) >= CITATION_TOP_MIN


def placeholder_idx(shape):
    try:
        if shape.is_placeholder:
            return shape.placeholder_format.idx
    except Exception:
        pass
    return None


def walk_shapes(shapes):
    """Yield every shape, descending into groups."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def effective_size(run, para):
    """Explicit run size, else the paragraph-level size (how the builders set fonts).
    None means fully inherited from the layout/master — not audited."""
    return run.font.size if run.font.size is not None else para.font.size


def font_floor_for(shape, level, in_table):
    if in_table:
        return FLOOR_TABLE, "table cell"
    if is_citation(shape):
        return FLOOR_CITATION, "citation"
    if is_badge(shape):
        return FLOOR_BADGE, "badge"
    idx = placeholder_idx(shape)
    if idx == 0:
        return None, "title"  # title size is template-controlled; not audited
    if idx is not None and idx >= 1:
        return (FLOOR_BODY, "body") if level == 0 else (FLOOR_SUBBULLET, "sub-bullet")
    return FLOOR_DIAGRAM, "diagram/textbox"


def extract_deck(pptx_path):
    """Extract per-slide text, font runs, and shape inventory from the built deck."""
    prs = Presentation(pptx_path)
    slide_area = prs.slide_width * prs.slide_height
    slides = []

    for s_i, slide in enumerate(prs.slides, start=1):
        info = {
            "number": s_i,
            "texts": [],          # content text fragments (used for numeric checks)
            "ref_texts": [],      # badge/citation boxes — reference elements, exempt
                                  # from orphan/added checks (authenticity is MANUAL)
            "font_violations": [],
            "pictures": [],       # (name, area_fraction)
            "editable_shapes": 0,
            "has_citation": False,
            "has_full_slide_image": False,
            "notes": "",
        }

        for shape in walk_shapes(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                area = (shape.width or 0) * (shape.height or 0)
                frac = area / slide_area if slide_area else 0
                info["pictures"].append((shape.name, frac))
                if frac > FULL_SLIDE_IMAGE_FRACTION:
                    info["has_full_slide_image"] = True
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue  # children are walked individually

            if getattr(shape, "has_table", False) and shape.has_table:
                info["editable_shapes"] += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        txt = cell.text_frame.text
                        if txt.strip():
                            info["texts"].append(txt)
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                sz = effective_size(run, para)
                                if sz is not None and sz.pt < FLOOR_TABLE:
                                    info["font_violations"].append(
                                        (f"{sz.pt:.0f}pt", FLOOR_TABLE, "table cell",
                                         run.text[:40]))
                continue

            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                txt = shape.text_frame.text
                if txt.strip():
                    info["editable_shapes"] += 1
                    if is_citation(shape) or is_badge(shape):
                        info["ref_texts"].append(txt)
                        if is_citation(shape):
                            info["has_citation"] = True
                    else:
                        info["texts"].append(txt)
                for para in shape.text_frame.paragraphs:
                    floor, ctx = font_floor_for(shape, para.level, in_table=False)
                    if floor is None:
                        continue
                    for run in para.runs:
                        sz = effective_size(run, para)
                        if sz is not None and sz.pt < floor:
                            info["font_violations"].append(
                                (f"{sz.pt:.0f}pt", floor, ctx, run.text[:40]))

        try:
            if slide.has_notes_slide:
                info["notes"] = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            pass

        slides.append(info)

    return slides


# ═══════════════════════════════════════════════════
#  Checks
# ═══════════════════════════════════════════════════

def check_claims_paper(claims_data, deck_slides):
    """Paper mode: each claim's numeric tokens must appear on its assigned slide."""
    results = []
    for claim in claims_data.get("claims", []):
        toks = [t.rstrip("%") for t in numeric_tokens(claim.get("value", ""))]
        slide_no = claim.get("slide")
        target = None
        if isinstance(slide_no, int) and 1 <= slide_no <= len(deck_slides):
            target = deck_slides[slide_no - 1]

        if not toks:
            results.append((claim, "⚠️", "claim has no numeric content to verify"))
            continue
        if target is None:
            # unassigned: search whole deck
            found_on = [
                s["number"] for s in deck_slides
                if all(t in [x.rstrip("%") for x in numeric_tokens(" ".join(s["texts"]))]
                       for t in toks)
            ]
            if found_on:
                results.append((claim, "⚠️",
                                f"no slide assigned; all numbers found on slide(s) "
                                f"{found_on} — assign and re-run"))
            else:
                results.append((claim, "❌", "no slide assigned and numbers not found "
                                             "anywhere in deck"))
            continue

        slide_toks = [t.rstrip("%") for t in numeric_tokens(" ".join(target["texts"]))]
        missing = [t for t in toks if t not in slide_toks]
        if not missing:
            results.append((claim, "✅", f"all numbers present on slide {slide_no}"))
        else:
            results.append((claim, "❌",
                            f"missing on slide {slide_no}: {', '.join(missing)}"))
    return results


def check_orphans(deck_slides, source_tokens, extra_known, skip_first_slide=True):
    """Numbers in the deck that trace to neither source nor claims are orphans."""
    known = set(source_tokens) | set(extra_known)
    known.add("95")  # 95% CI boilerplate
    orphans = []
    for s in deck_slides:
        if skip_first_slide and s["number"] == 1:
            continue  # title slide carries presenter/date metadata, not source data
        for tok in numeric_tokens(" ".join(s["texts"])):
            base = tok.rstrip("%")
            if base not in known and tok not in known:
                orphans.append((s["number"], tok))
    # dedupe, keep order
    seen, out = set(), []
    for o in orphans:
        if o not in seen:
            seen.add(o)
            out.append(o)
    return out


def check_placeholders(deck_slides):
    hits = []
    for s in deck_slides:
        for txt in s["texts"] + s["ref_texts"]:
            for line in txt.splitlines():
                if "[CHECK" in line.upper():
                    hits.append((s["number"], line.strip()[:80]))
    return hits


def _is_ref_element(el):
    """Source-deck badge/citation boxes — same positional exemption as the deck side."""
    if el.get("role") == "badge":
        return True
    pos = el.get("position") or {}
    left, top = pos.get("left") or 0, pos.get("top") or 0
    return (left >= BADGE_LEFT_MIN and top <= BADGE_TOP_MAX) or \
           (left >= CITATION_LEFT_MIN and top >= CITATION_TOP_MIN)


def check_pptx_mode(claims_data, deck_slides):
    """pptx mode: per-slide two-way numeric comparison, source deck = ground truth."""
    src_slides = claims_data.get("slides", [])
    results = []
    n = min(len(src_slides), len(deck_slides))
    if len(src_slides) != len(deck_slides):
        results.append(("slide count", "❌",
                        f"source {len(src_slides)} slides vs output {len(deck_slides)}"))
    for i in range(n):
        src_text = " ".join(
            el.get("full_text", "") or " ".join(str(c) for row in el.get("rows", []) for c in row)
            for el in src_slides[i].get("elements", [])
            if not _is_ref_element(el)
        )
        src_toks = set(t.rstrip("%") for t in numeric_tokens(src_text))
        out_toks = set(t.rstrip("%") for t in
                       numeric_tokens(" ".join(deck_slides[i]["texts"])))
        missing = sorted(src_toks - out_toks)
        added = sorted(out_toks - src_toks)
        label = f"slide {i + 1}"
        if not missing and not added:
            results.append((label, "✅", "numeric content matches source slide"))
        else:
            parts = []
            if missing:
                parts.append(f"missing from output: {', '.join(missing)}")
            if added:
                parts.append(f"not in source: {', '.join(added)}")
            # missing numbers may be legitimate (content trimmed to meet font floors,
            # rasterized charts) — flag, human decides; added numbers are hard fails
            verdict = "❌" if added else "⚠️"
            results.append((label, verdict, "; ".join(parts)))
    return results


# ═══════════════════════════════════════════════════
#  Report
# ═══════════════════════════════════════════════════

MANUAL_LINES = [
    "Treatment-arm labels correct and not swapped (check labels against source, not just numbers)",
    "Endpoints correctly tagged primary vs secondary",
    "Biomarker cutoffs/thresholds correct in context (CLDN18.2, HER2, MSI-H, ctDNA, PD-L1 CPS/TPS, ...)",
    "Drug names, doses, schedules read correctly in context",
    "Every citation traces to a real reference in the source (no fabricated references)",
    "Figures embedded at correct aspect ratio; no overflow/cut-off text (eyeball rendered slides)",
    "Moffitt branding present on every slide (template inheritance)",
]


def build_report(deck_path, args, deck_slides, claim_results, orphans, placeholders,
                 mode):
    lines = []
    counts = {"✅": 0, "❌": 0, "⚠️": 0}

    def add(verdict, text):
        counts[verdict] += 1
        lines.append(f"- {verdict} {text}")

    today = datetime.date.today().isoformat()
    hdr = [
        f"# QA Report — {os.path.basename(deck_path)}",
        "",
        f"Deck: `{os.path.basename(deck_path)}` | Source: "
        f"`{os.path.basename(args.claims)}`"
        + (f" + `{os.path.basename(args.source_text)}`" if args.source_text else "")
        + f" | Built: {today} | Mode: {mode}",
        "",
        "Reviewer: ______________  Date: ______________  Verdict: ______________",
        "",
        "## 1. Data accuracy",
        "",
    ]

    for item, verdict, msg in claim_results:
        if isinstance(item, dict):  # paper-mode claim
            add(verdict, f"[{item.get('category', 'other')}] {item.get('label', item.get('id', '?'))}: "
                         f"`{item.get('value', '')}` — {msg}"
                         + (f" (source: {item['source_hint']})" if item.get("source_hint") else ""))
        else:  # pptx-mode per-slide line
            add(verdict, f"{item}: {msg}")

    if args.source_text or mode == "pptx":
        if orphans:
            for slide_no, tok in orphans:
                add("❌", f"orphan number `{tok}` on slide {slide_no} — not found in source; "
                          f"verify or remove")
        else:
            add("✅", "no orphan numbers — every number in the deck traces to the source")
    else:
        add("⚠️", "orphan-number check SKIPPED (no --source-text provided)")

    if placeholders:
        for slide_no, txt in placeholders:
            add("❌", f"unresolved [CHECK] placeholder on slide {slide_no}: “{txt}”")
    else:
        add("✅", "no unresolved [CHECK] placeholders")

    lines.append("")
    lines.append("## 2. Legibility (font floors)")
    lines.append("")
    any_font = False
    for s in deck_slides:
        for got, floor, ctx, snippet in s["font_violations"]:
            any_font = True
            add("❌", f"slide {s['number']}: {ctx} run at {got} (floor {floor}pt): “{snippet}”")
    if not any_font:
        add("✅", "no text below font-size floors "
                  "(body 18 / sub-bullet 16 / table 14 / diagram 12; citation 10, badge 11 exempt)")

    lines.append("")
    lines.append("## 3. Editability")
    lines.append("")
    any_edit = False
    for s in deck_slides:
        if s["has_full_slide_image"]:
            any_edit = True
            add("❌", f"slide {s['number']} is dominated by a single full-slide image — "
                      f"violates the editability principle")
    for s in deck_slides:
        for name, frac in s["pictures"]:
            if not s["has_full_slide_image"] and frac > 0:
                add("⚠️", f"slide {s['number']}: picture `{name}` covers {frac:.0%} of slide — "
                          f"confirm it is a source figure or a justified rasterization "
                          f"(chart / complex group), not flattened content")
                any_edit = True
    if not any_edit:
        add("✅", "no pictures needing review; all content shapes are editable objects")

    lines.append("")
    lines.append("## 4. Structure & citations")
    lines.append("")
    missing_cite = [s["number"] for s in deck_slides[1:] if not s["has_citation"]]
    if missing_cite:
        add("⚠️", f"no citation textbox detected on content slide(s) {missing_cite} "
                  f"(recognized by bottom-right position)")
    else:
        add("✅", "citation textbox present on every content slide")

    lines.append("")
    lines.append("## 5. Manual review (script cannot verify — check each)")
    lines.append("")
    for m in MANUAL_LINES:
        lines.append(f"- [ ] MANUAL: {m}")

    summary = (f"**Summary:** {counts['✅']} ✅ | {counts['❌']} ❌ | {counts['⚠️']} ⚠️ | "
               f"{len(MANUAL_LINES)} manual items. "
               + ("**Deck is NOT final — fix every ❌ and re-run.**" if counts["❌"]
                  else "No hard failures — complete the manual items to finalize."))
    report = "\n".join(hdr + lines + ["", summary, ""])
    return report, counts


def write_notes(pptx_path, deck_slides, counts):
    """Optionally stamp a one-line QA status into each slide's speaker notes."""
    prs = Presentation(pptx_path)
    stamp = (f"[QA {datetime.date.today().isoformat()}: "
             f"{counts['❌']} fail / {counts['⚠️']} warn — see QA report]")
    for slide, info in zip(prs.slides, deck_slides):
        tf = slide.notes_slide.notes_text_frame
        existing = tf.text or ""
        if stamp not in existing:
            tf.text = (existing + "\n" + stamp).strip()
    prs.save(pptx_path)


def main():
    ap = argparse.ArgumentParser(description="Deterministic deck QA cross-check")
    ap.add_argument("deck", help="Built .pptx to verify")
    ap.add_argument("--claims", required=True,
                    help="paper mode: claims JSON; pptx mode: parse_pptx.py JSON")
    ap.add_argument("--mode", choices=["paper", "pptx"], default="paper")
    ap.add_argument("--source-text", default=None,
                    help="paper mode: paper markdown/text for orphan-number check")
    ap.add_argument("--output", "-o", default=None,
                    help="QA report path (default: <deck>_QA.md)")
    ap.add_argument("--notes", action="store_true",
                    help="also stamp QA status into speaker notes")
    args = ap.parse_args()

    for path in [args.deck, args.claims] + ([args.source_text] if args.source_text else []):
        if not os.path.exists(path):
            print(f"Error: not found: {path}", file=sys.stderr)
            sys.exit(2)

    with open(args.claims) as f:
        claims_data = json.load(f)

    deck_slides = extract_deck(args.deck)

    if args.mode == "paper":
        claim_results = check_claims_paper(claims_data, deck_slides)
        claim_tokens = [t.rstrip("%") for c in claims_data.get("claims", [])
                        for t in numeric_tokens(c.get("value", ""))]
        meta_tokens = numeric_tokens(" ".join(str(v) for v in
                                              claims_data.get("deck", {}).values()))
        source_tokens = []
        if args.source_text:
            with open(args.source_text, errors="replace") as f:
                source_tokens = [t.rstrip("%") for t in numeric_tokens(f.read())]
            orphans = check_orphans(deck_slides, source_tokens,
                                    claim_tokens + meta_tokens)
        else:
            orphans = []
    else:
        claim_results = check_pptx_mode(claims_data, deck_slides)
        src_text = " ".join(
            el.get("full_text", "") or " ".join(str(c) for row in el.get("rows", []) for c in row)
            for s in claims_data.get("slides", []) for el in s.get("elements", [])
        )
        orphans = check_orphans(deck_slides,
                                [t.rstrip("%") for t in numeric_tokens(src_text)], [])

    placeholders = check_placeholders(deck_slides)

    report, counts = build_report(args.deck, args, deck_slides, claim_results,
                                  orphans, placeholders, args.mode)

    out_path = args.output or os.path.splitext(args.deck)[0] + "_QA.md"
    with open(out_path, "w") as f:
        f.write(report)
    print(f"QA report: {os.path.abspath(out_path)}")
    print(f"Result: {counts['✅']} pass / {counts['❌']} fail / {counts['⚠️']} warn")

    if args.notes:
        write_notes(args.deck, deck_slides, counts)
        print("Speaker-notes QA stamps written.")

    sys.exit(1 if counts["❌"] else 0)


if __name__ == "__main__":
    main()
