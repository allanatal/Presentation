import os, sys, json, copy
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import build_deck as bd

CLAIMS = {
    "deck": {"title": "T", "source": "p.pdf", "nct": "NCT00000000", "date": "2026-07-12"},
    "claims": [
        {"id": "eff-01", "category": "efficacy", "label": "Median OS IP",
         "value": "19.4 months (95% CI 17.1-22.9)", "slide": None, "source_hint": "Fig 2A"},
        {"id": "eff-03", "category": "efficacy", "label": "OS HR",
         "value": "HR 0.67 (95% CI 0.50-0.90); P = .01", "slide": None, "source_hint": "Fig 2A"},
        {"id": "dem-01", "category": "demographics", "label": "Age",
         "value": "IP 60 (24-70); PS 56 (23-74)", "slide": None, "source_hint": "Table 1"},
    ],
}


# ═══════════ Task 1: Claims core ═══════════

def test_value_lookup():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    assert c.value("eff-01") == "19.4 months (95% CI 17.1-22.9)"


def test_unknown_claim_raises():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    with pytest.raises(KeyError):
        c.value("nope-99")


def test_substitute_expands_and_assigns():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    out = c.substitute("Median OS {{eff-01}}; {{eff-03}}", slide_no=10)
    assert out == "Median OS 19.4 months (95% CI 17.1-22.9); HR 0.67 (95% CI 0.50-0.90); P = .01"
    assert c.referenced["eff-01"] == 10
    assert c.referenced["eff-03"] == 10


def test_substitute_unknown_token_raises():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    with pytest.raises(KeyError):
        c.substitute("bad {{nope-99}}", slide_no=1)


def test_explicit_assign():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    c.assign("dem-01", 8)
    assert c.referenced["dem-01"] == 8


def test_first_assignment_wins():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    c.assign("eff-01", 5)
    c.assign("eff-01", 9)
    assert c.referenced["eff-01"] == 5


def test_resolved_json_fills_slide_and_preserves_original():
    data = copy.deepcopy(CLAIMS)
    c = bd.Claims(data)
    c.assign("eff-01", 10)
    c.assign("dem-01", 8)
    resolved = c.resolved_json()
    by = {x["id"]: x for x in resolved["claims"]}
    assert by["eff-01"]["slide"] == 10
    assert by["dem-01"]["slide"] == 8
    assert by["eff-03"]["slide"] is None
    assert data["claims"][0]["slide"] is None  # original untouched


def test_unreferenced_lists_unused_ids():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    c.assign("eff-01", 10)
    assert set(c.unreferenced()) == {"eff-03", "dem-01"}


# ═══════════ Task 2: recursive spec substitution ═══════════

def test_substitute_in_place_walks_nested_structures():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    sdata = {
        "type": "figure",
        "key_message": "OS {{eff-01}}",
        "caption": ["deaths line", "HR is {{eff-03}}"],
        "rows": [{"label": "Age", "vals": ["{{dem-01}}"]}],
    }
    bd.substitute_in_place(sdata, c, slide_no=10)
    assert sdata["key_message"] == "OS 19.4 months (95% CI 17.1-22.9)"
    assert sdata["caption"][1] == "HR is HR 0.67 (95% CI 0.50-0.90); P = .01"
    assert sdata["rows"][0]["vals"][0] == "IP 60 (24-70); PS 56 (23-74)"
    assert c.referenced == {"eff-01": 10, "eff-03": 10, "dem-01": 10}


def test_slide_level_claims_assigned():
    c = bd.Claims(copy.deepcopy(CLAIMS))
    sdata = {"type": "table", "claims": ["dem-01"], "rows": []}
    bd.assign_slide_claims(sdata, c, slide_no=8)
    assert c.referenced["dem-01"] == 8


# ═══════════ Deck-build fixtures ═══════════

import shutil
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _skill_root():
    # scripts/ -> skill root
    return os.path.dirname(os.path.dirname(os.path.abspath(bd.__file__)))


def _new_ctx(tmp_path):
    out = str(tmp_path / "out.pptx")
    shutil.copy(os.path.join(_skill_root(), "references", "template.pptx"), out)
    prs = Presentation(out)
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}
    ctx = bd.Ctx(prs=prs, layouts=layouts, study="TEST-01",
                 citation="Author et al. 2026", skill_root=_skill_root(),
                 images_base=str(tmp_path))
    return ctx, out


def _texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
    return "\n".join(out)


# ═══════════ Task 3: title / bullets / figure renderers ═══════════

def test_render_title_sets_trial_name(tmp_path):
    ctx, out = _new_ctx(tmp_path)
    s = {"type": "title", "title": "TEST-01: A Trial", "subtitle": "Phase 3",
         "authors": "Author A", "affiliation": "Center X", "date": "Journal 2026"}
    bd.render_title(ctx, s, 1)
    slide = ctx.prs.slides[0]
    assert "TEST-01: A Trial" in _texts(slide)
    assert "Phase 3" in _texts(slide)


def test_render_bullets_counts_and_badge(tmp_path):
    ctx, out = _new_ctx(tmp_path)
    s = {"type": "bullets", "title": "Background",
         "bullets": [{"text": "Point one"}, {"text": "Point two", "sub": ["sub a"]}]}
    bd.render_bullets(ctx, s, 2)
    slide = ctx.prs.slides[0]
    txt = _texts(slide)
    assert "Point one" in txt and "Point two" in txt and "sub a" in txt
    assert "TEST-01" in txt              # badge
    assert "Author et al. 2026" in txt   # citation


def test_render_bullets_font_floors(tmp_path):
    ctx, out = _new_ctx(tmp_path)
    s = {"type": "bullets", "title": "T",
         "bullets": [{"text": "L0"}, {"text": "L0b", "sub": ["L1"]}]}
    bd.render_bullets(ctx, s, 2)
    slide = ctx.prs.slides[0]
    body = None
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 1:
            body = sh
    sizes = {}
    for p in body.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                sizes.setdefault(p.level, set()).add(r.font.size.pt)
    assert all(x >= 18 for x in sizes.get(0, {18}))
    assert all(x >= 16 for x in sizes.get(1, {16}))


def test_render_figure_embeds_picture(tmp_path):
    from PIL import Image as PILImage
    p = tmp_path / "fig.png"
    PILImage.new("RGB", (800, 600), "white").save(p)
    ctx, out = _new_ctx(tmp_path)
    s = {"type": "figure", "title": "OS", "image": "fig.png",
         "key_message": "median", "caption": ["line1"]}
    bd.render_figure(ctx, s, 3)
    slide = ctx.prs.slides[0]
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes)
