#!/usr/bin/env python3
"""
Build a Moffitt-styled PPTX from parsed slide data (output of parse_pptx.py).

Usage:
    python build_from_parsed.py parsed.json [--output output.pptx] \
        [--study-name "Study Name"] [--citation "Author et al. 2026"]

This script reads the structured JSON from parse_pptx.py and rebuilds
every slide using the Moffitt template and shared slide-builder patterns.
"""

import argparse
import json
import os
import shutil
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

try:
    from PIL import Image
except ImportError:
    Image = None

# ── Shared paths ──
TEMPLATE = "/mnt/skills/user/academic-paper-to-pptx/references/template.pptx"
LOGO_PATH = "/mnt/skills/user/academic-paper-to-pptx/references/Picture_3.x-wmf"

# ── Colors (from style-spec.md) ──
TITLE_BLUE = RGBColor(0x00, 0x33, 0x66)
RED = RGBColor(0xFF, 0x00, 0x00)
BODY = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x99, 0x99, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
LIGHT_BG = RGBColor(0xF0, 0xF4, 0xF8)
ARM_TEAL = RGBColor(0x2E, 0x7D, 0x7D)
ARM_BLUE = RGBColor(0x44, 0x72, 0xC4)

FONT = "Arial"
SLIDE_WIDTH = 12192000  # EMU


# ═══════════════════════════════════════════════════
#  Utility functions
# ═══════════════════════════════════════════════════

def add_badge_and_citation(slide, study_name="", citation_text=""):
    """Add study name badge (top-right) and citation (bottom-right)."""
    if study_name:
        badge = slide.shapes.add_textbox(
            Emu(9984826), Emu(157655), Emu(1671098), Emu(369332)
        )
        p = badge.text_frame.paragraphs[0]
        p.text = study_name
        p.font.size = Pt(11)
        p.font.color.rgb = TITLE_BLUE
        p.alignment = PP_ALIGN.RIGHT

    if citation_text:
        cite = slide.shapes.add_textbox(
            Emu(9987101), Emu(6463328), Emu(2204899), Emu(307777)
        )
        p = cite.text_frame.paragraphs[0]
        p.text = citation_text
        p.font.size = Pt(10)
        p.font.color.rgb = GRAY
        p.alignment = PP_ALIGN.RIGHT


def get_title_text(slide_data):
    """Extract the title text from a parsed slide."""
    for el in slide_data.get("elements", []):
        if el.get("type") == "text" and el.get("role") == "title":
            return el.get("full_text", "").strip()
    return ""


def get_subtitle_text(slide_data):
    """Extract subtitle text from a parsed slide."""
    for el in slide_data.get("elements", []):
        if el.get("type") == "text" and el.get("role") == "subtitle":
            return el.get("full_text", "").strip()
    return ""


def get_body_paragraphs(slide_data):
    """Collect all body paragraphs from a parsed slide, flattened."""
    paragraphs = []
    for el in slide_data.get("elements", []):
        if el.get("type") == "text" and el.get("role") == "body":
            for p in el.get("paragraphs", []):
                if p.get("text", "").strip():
                    paragraphs.append(p)
    return paragraphs


def get_footnote_text(slide_data):
    """Extract footnote/caption text from a parsed slide."""
    footnotes = []
    for el in slide_data.get("elements", []):
        if el.get("type") == "text" and el.get("role") == "footnote":
            txt = el.get("full_text", "").strip()
            if txt:
                footnotes.append(txt)
    return " | ".join(footnotes) if footnotes else ""


def get_table_element(slide_data):
    """Get the first table element from a parsed slide."""
    for el in slide_data.get("elements", []):
        if el.get("type") == "table":
            return el
    return None


def get_image_elements(slide_data):
    """Get all image elements from a parsed slide."""
    images = []
    for el in slide_data.get("elements", []):
        if el.get("type") == "image" and el.get("image_path"):
            if os.path.exists(el["image_path"]):
                images.append(el)
    return images


def get_chart_image_path(slide_data, slide_idx, images_dir):
    """Get the rasterized chart image path for a chart slide."""
    # Check if a pre-rasterized chart image exists
    for suffix in ["_chart.jpg", "_chart.png"]:
        candidate = os.path.join(images_dir, f"chart_slide-{slide_idx + 1:02d}{suffix}")
        if os.path.exists(candidate):
            return candidate
    # Fall back to full slide raster
    candidate = os.path.join(images_dir, f"slide_raster-{slide_idx + 1:02d}.jpg")
    if os.path.exists(candidate):
        return candidate
    return None


# ═══════════════════════════════════════════════════
#  Slide builders — one per slide type
# ═══════════════════════════════════════════════════

def build_title_slide(prs, layouts, slide_data, study_name, citation):
    """Build the opening title slide."""
    slide = prs.slides.add_slide(layouts["Title Slide"])

    title_text = get_title_text(slide_data)
    subtitle_text = get_subtitle_text(slide_data)

    # Also check body elements for subtitle-like content on title slides
    if not subtitle_text:
        body_paras = get_body_paragraphs(slide_data)
        if body_paras:
            subtitle_text = "\n".join(p["text"] for p in body_paras[:3])

    # Set title
    title_ph = slide.placeholders[0]
    title_ph.text = title_text
    for para in title_ph.text_frame.paragraphs:
        para.font.size = Pt(22)
        para.font.bold = True
        para.font.color.rgb = RED
        para.font.name = FONT

    # Set subtitle
    if subtitle_text and 1 in slide.placeholders:
        sub_ph = slide.placeholders[1]
        sub_ph.text = subtitle_text
        for para in sub_ph.text_frame.paragraphs:
            para.font.size = Pt(18)
            para.font.bold = True
            para.font.color.rgb = BODY
            para.font.name = FONT

    # Add logo if available
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(
                LOGO_PATH,
                Emu(4790362), Emu(798667),
                Emu(2266391), Emu(578214),
            )
        except Exception:
            pass  # Logo is optional

    return slide


def build_content_slide(prs, layouts, slide_data, study_name, citation):
    """Build a Title and Content slide with bullets."""
    slide = prs.slides.add_slide(layouts["Title and Content"])

    # Title
    title_text = get_title_text(slide_data)
    if title_text:
        slide.placeholders[0].text = title_text

    # Body
    body_paras = get_body_paragraphs(slide_data)
    if body_paras and 1 in slide.placeholders:
        tf = slide.placeholders[1].text_frame
        tf.clear()

        first = True
        for para in body_paras:
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()

            p.text = para["text"]
            p.level = min(para.get("level", 0), 1)
            p.font.name = FONT

            # Apply font size minimums
            if p.level == 0:
                p.font.size = Pt(18)
            else:
                p.font.size = Pt(16)

            if para.get("bold"):
                p.font.bold = True

    add_badge_and_citation(slide, study_name, citation)
    return slide


def build_table_slide(prs, layouts, slide_data, study_name, citation):
    """Build a slide with a centered table."""
    slide = prs.slides.add_slide(layouts["Title Only"])

    # Title
    title_text = get_title_text(slide_data)
    if title_text:
        slide.placeholders[0].text = title_text

    table_el = get_table_element(slide_data)
    if not table_el:
        add_badge_and_citation(slide, study_name, citation)
        return slide

    rows = table_el["rows"]
    n_rows = len(rows)
    n_cols = table_el.get("n_cols", len(rows[0]) if rows else 0)

    if n_rows == 0 or n_cols == 0:
        add_badge_and_citation(slide, study_name, citation)
        return slide

    # Trim rows if needed to fit at 14pt
    max_rows = 14
    trimmed = False
    if n_rows > max_rows:
        # Keep header + first (max_rows - 1) data rows
        rows = rows[:max_rows]
        n_rows = max_rows
        trimmed = True

    # Calculate column widths and center
    available_w = Inches(11.5)
    col_w = int(available_w / n_cols)
    total_w = col_w * n_cols
    table_x = (SLIDE_WIDTH - total_w) // 2

    row_h = Emu(320000)
    tbl_shape = slide.shapes.add_table(
        n_rows, n_cols,
        table_x, Emu(1900000),
        total_w, row_h * n_rows,
    )
    tbl = tbl_shape.table

    for i in range(n_cols):
        tbl.columns[i].width = col_w

    for r, row in enumerate(rows):
        for c in range(min(len(row), n_cols)):
            cell = tbl.cell(r, c)
            cell.text = str(row[c])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(14)  # MINIMUM 14pt
                p.font.name = FONT
                if r == 0:
                    p.font.bold = True
                    p.font.color.rgb = WHITE
                    p.alignment = PP_ALIGN.CENTER
                else:
                    p.font.color.rgb = BODY
                    if c >= 1:
                        p.alignment = PP_ALIGN.CENTER

            # Header row fill
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY

    # Trimming note
    if trimmed:
        note_y = Emu(1900000) + row_h * n_rows + Emu(100000)
        note = slide.shapes.add_textbox(table_x, note_y, total_w, Emu(300000))
        p = note.text_frame.paragraphs[0]
        p.text = "Selected data shown; see original for full table."
        p.font.size = Pt(12)
        p.font.italic = True
        p.font.color.rgb = GRAY
        p.font.name = FONT

    add_badge_and_citation(slide, study_name, citation)
    return slide


def build_figure_slide(prs, layouts, slide_data, study_name, citation):
    """Build a slide with embedded image(s)."""
    slide = prs.slides.add_slide(layouts["Title Only"])

    # Title
    title_text = get_title_text(slide_data)
    if title_text:
        slide.placeholders[0].text = title_text

    # Key message from subtitle
    subtitle = get_subtitle_text(slide_data)
    msg_offset = 0
    if subtitle:
        msg = slide.shapes.add_textbox(
            Emu(838200), Emu(1400000), Emu(10515600), Emu(400000)
        )
        msg.text_frame.word_wrap = True
        p = msg.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.italic = True
        p.font.color.rgb = TITLE_BLUE
        p.font.name = FONT
        msg_offset = 150000

    # Embed images
    img_els = get_image_elements(slide_data)
    if img_els and Image:
        img_path = img_els[0]["image_path"]
        _embed_image(slide, img_path, y_offset=msg_offset)

    add_badge_and_citation(slide, study_name, citation)
    return slide


def build_chart_slide(prs, layouts, slide_data, slide_idx, images_dir, study_name, citation):
    """Build a slide with a rasterized chart image."""
    slide = prs.slides.add_slide(layouts["Title Only"])

    # Title
    title_text = get_title_text(slide_data)
    if title_text:
        slide.placeholders[0].text = title_text

    # Try to find the rasterized chart
    chart_path = get_chart_image_path(slide_data, slide_idx, images_dir)
    if chart_path and Image:
        _embed_image(slide, chart_path)
    else:
        # Fallback: add a note that the chart needs manual insertion
        note = slide.shapes.add_textbox(
            Emu(2000000), Emu(3000000), Emu(8000000), Emu(1000000)
        )
        note.text_frame.word_wrap = True
        p = note.text_frame.paragraphs[0]
        p.text = "[Chart from source slide — rasterize source and embed manually]"
        p.font.size = Pt(16)
        p.font.italic = True
        p.font.color.rgb = GRAY
        p.font.name = FONT
        p.alignment = PP_ALIGN.CENTER

    add_badge_and_citation(slide, study_name, citation)
    return slide


def build_diagram_slide(prs, layouts, slide_data, slide_idx, images_dir, study_name, citation):
    """Build a slide from a rasterized diagram (grouped shapes)."""
    # Diagrams are complex grouped shapes — safest to rasterize
    return build_chart_slide(prs, layouts, slide_data, slide_idx, images_dir, study_name, citation)


def _embed_image(slide, img_path, y_offset=0):
    """Embed an image centered on the slide with proper aspect ratio."""
    if not Image or not os.path.exists(img_path):
        return

    img = Image.open(img_path)
    img_w, img_h = img.size
    ratio = img_w / img_h

    max_w = Emu(10500000)
    max_h = Emu(4300000)
    fig_y = Emu(1900000 + y_offset)

    w = max_w
    h = int(w / ratio)
    if h > max_h:
        h = max_h
        w = int(h * ratio)

    fig_x = Emu(838200) + (max_w - w) // 2
    slide.shapes.add_picture(img_path, fig_x, fig_y, w, h)


# ═══════════════════════════════════════════════════
#  Speaker notes transfer
# ═══════════════════════════════════════════════════

def transfer_notes(output_slide, slide_data):
    """Copy speaker notes from source to output slide."""
    if slide_data.get("has_notes") and slide_data.get("notes_text"):
        try:
            notes_slide = output_slide.notes_slide
            notes_slide.notes_text_frame.text = slide_data["notes_text"]
        except Exception:
            pass  # Notes are optional; don't fail the build


# ═══════════════════════════════════════════════════
#  Main build orchestrator
# ═══════════════════════════════════════════════════

def build_presentation(parsed_data, output_path, study_name="", citation="", images_dir="/home/claude/source_images"):
    """Build the complete Moffitt-styled presentation from parsed data."""
    # Copy template
    shutil.copy(TEMPLATE, output_path)
    prs = Presentation(output_path)

    master = prs.slide_masters[0]
    layouts = {l.name: l for l in master.slide_layouts}

    # Verify required layouts exist
    required = ["Title Slide", "Title and Content", "Title Only"]
    for name in required:
        if name not in layouts:
            print(f"WARNING: Layout '{name}' not found in template. Available: {list(layouts.keys())}", file=sys.stderr)

    slides = parsed_data.get("slides", [])
    print(f"Building {len(slides)} slides...", file=sys.stderr)

    for slide_data in slides:
        slide_type = slide_data.get("slide_type", "content_slide")
        slide_idx = slide_data.get("index", 0)

        try:
            if slide_type == "title_slide":
                out_slide = build_title_slide(prs, layouts, slide_data, study_name, citation)
            elif slide_type == "table_slide":
                out_slide = build_table_slide(prs, layouts, slide_data, study_name, citation)
            elif slide_type == "figure_slide":
                out_slide = build_figure_slide(prs, layouts, slide_data, study_name, citation)
            elif slide_type == "chart_slide":
                out_slide = build_chart_slide(prs, layouts, slide_data, slide_idx, images_dir, study_name, citation)
            elif slide_type == "diagram_slide":
                out_slide = build_diagram_slide(prs, layouts, slide_data, slide_idx, images_dir, study_name, citation)
            else:
                out_slide = build_content_slide(prs, layouts, slide_data, study_name, citation)

            # Transfer speaker notes
            transfer_notes(out_slide, slide_data)

            title = get_title_text(slide_data) or "(no title)"
            print(f"  ✓ Slide {slide_idx}: {slide_type:15s} → {title[:50]}", file=sys.stderr)

        except Exception as e:
            print(f"  ✗ Slide {slide_idx}: {slide_type} FAILED — {e}", file=sys.stderr)
            # Build a placeholder slide so slide count is preserved
            fallback = prs.slides.add_slide(layouts["Title Only"])
            fallback.placeholders[0].text = get_title_text(slide_data) or f"Slide {slide_idx + 1}"
            note = fallback.shapes.add_textbox(
                Emu(2000000), Emu(3000000), Emu(8000000), Emu(1000000)
            )
            p = note.text_frame.paragraphs[0]
            p.text = f"[Conversion error: {e}]"
            p.font.size = Pt(14)
            p.font.color.rgb = GRAY
            p.font.name = FONT
            p.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"\n✅ Saved {len(slides)} slides to {output_path}", file=sys.stderr)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Build Moffitt PPTX from parsed JSON")
    parser.add_argument("json_path", help="Path to parsed JSON (from parse_pptx.py)")
    parser.add_argument("--output", "-o", default="/home/claude/output.pptx", help="Output PPTX path")
    parser.add_argument("--study-name", default="", help="Study name for badge")
    parser.add_argument("--citation", default="", help="Citation text for footer")
    parser.add_argument("--images-dir", default="/home/claude/source_images", help="Directory with extracted/rasterized images")
    args = parser.parse_args()

    if not os.path.exists(args.json_path):
        print(f"Error: File not found: {args.json_path}", file=sys.stderr)
        sys.exit(1)

    with open(args.json_path) as f:
        parsed_data = json.load(f)

    build_presentation(
        parsed_data,
        args.output,
        study_name=args.study_name,
        citation=args.citation,
        images_dir=args.images_dir,
    )


if __name__ == "__main__":
    main()
